#!/usr/bin/env python3
"""Offline checks for Family Book v2.1.

These use a temporary folder as the user's computer. They do not call Excel,
and they do not read or write a real Family Book folder. Receipt files written
here start with ``TEST - ``.

    python3 tests/v2.1/check_v21.py

The v2.1 cases need openpyxl. The v2.0 text checks imported at the end need
nothing else. PDF text extraction needs ``pypdf`` and is skipped when it is
missing. Regenerating the v2.0 fixture PDFs (``generate()``) needs
``reportlab``. Install both with ``pip install pypdf reportlab`` when you
want those checks. Excel is not required.
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import zipfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from bot.local.bookio import (  # noqa: E402
    advice,
    answer_question,
    apply_receipts,
    cards_empty,
    init_local_book,
    is_locked,
    mark_statement_added,
    migrate_v21_settings,
    read_settings,
    receipt_key,
    receipt_stats,
    record_report_sent,
    report_already_sent,
    report_totals,
    routine_gate,
    set_offer,
    short_store,
    suggest_statement,
    transactions_empty,
    week_period,
)
from bot.local.queue import LocalComputer, ReceiptQueue, drain, resolve_question  # noqa: E402


def _load_v2_identity():
    import importlib.util
    path = ROOT / "tests" / "v2.0" / "familybook_v2.py"
    spec = importlib.util.spec_from_file_location("familybook_v2", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules["familybook_v2"] = module
    spec.loader.exec_module(module)
    return module


V2 = _load_v2_identity()


class Fail(RuntimeError):
    pass


def check(condition, detail) -> None:
    if not condition:
        raise Fail(detail)


def receipt(store="Costco", when="2026-10-03 12:00", total=12.34, **extra) -> dict:
    body = extra.pop("body", None)
    if body is None:
        body = f"{store}\n{when}\nTotal {total:.2f}\nPaid cash\n"
    item = {
        "store": store,
        "when": when,
        "total": total,
        "number": extra.pop("number", ""),
        "last4": extra.pop("last4", ""),
        "category": extra.pop("category", "Meals"),
        "source": extra.pop("source", "photo"),
        "ext": extra.pop("ext", ".txt"),
        "body": body,
    }
    item.update(extra)
    return item


def new_book(folder: Path, year: int = 2026) -> Path:
    name = "TEST - Family Book - 2026.xlsx"
    return init_local_book(ROOT / "template" / "family-book-template.xlsx", folder, year, name)


def grid_amount(path: Path, sheet: str, coord: str):
    wb = load_workbook(path)
    return wb[sheet][coord].value


def snapshot_numbers(path: Path) -> dict:
    wb = load_workbook(path)
    found = {}
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    found[(name, cell.coordinate)] = cell.value
                elif isinstance(cell.value, str) and cell.value.startswith("="):
                    found[(name, cell.coordinate)] = cell.value
    return found


def test_keys_match_cloud() -> None:
    when = date(2026, 10, 3)
    check(short_store("Costco.ca") == "Costco", short_store("Costco.ca"))
    check(short_store("HOME DEPOT #7016") == "Home Depot", short_store("HOME DEPOT #7016"))
    blank = receipt_key("Costco Wholesale", "", when, 12.34)
    numbered = receipt_key("Costco", "#48291", when, 12.34)
    check(blank == "Costco|2026-10-03|12.34", blank)
    check(numbered == "Costco|48291", numbered)
    check(blank == V2.receipt_identity("Costco", "", when, 12.34), blank)
    check(numbered == V2.receipt_identity("Costco", "48291", when, 12.34), numbered)


def test_receipts_only_and_no_excel(tmp: Path) -> None:
    book = new_book(tmp)
    source = (ROOT / "bot" / "local" / "bookio.py").read_text(encoding="utf-8")
    check("subprocess" not in source, "book writer must not shell out")
    check("startfile" not in source and "osascript" not in source, "book writer must not launch an app")
    result = apply_receipts(
        book,
        [receipt()],
        tmp / "receipts",
        tmp / "backups",
        name_prefix="TEST - ",
        stamp="20261024-120000",
    )
    check(result["wrote"] == ["Costco|2026-10-03|12.34"], result)
    check(result["backup"], result)
    check(grid_amount(book, "Oct", "O6") == 12.34, grid_amount(book, "Oct", "O6"))
    check(transactions_empty(book), "transactions should stay empty")
    check(cards_empty(book), "cards should stay empty")
    check(routine_gate(book)["may_remind"] is False, routine_gate(book))
    check(routine_gate(book)["may_nudge"] is False, "nudge")
    check(routine_gate(book)["may_bank_income"] is False, "income")
    check(routine_gate(book)["may_bank_fixed"] is False, "fixed")
    wb = load_workbook(book)
    check(wb["Annual Summary"]["C47"].value == "off", wb["Annual Summary"]["C47"].value)
    check(wb["Annual Summary"]["C46"].value == "v2.1", wb["Annual Summary"]["C46"].value)
    check(read_settings(wb)["started"] == date.today().isoformat(), read_settings(wb))
    names = list((tmp / "receipts").iterdir())
    check(len(names) == 1 and names[0].name.startswith("TEST - "), names)
    check(names[0].name == "TEST - 2026-10-03 1200 - Costco - 12.34.txt", names[0].name)
    # The file is a real xlsx package. Opening it does not require Excel.
    package = zipfile.ZipFile(book)
    check("[Content_Types].xml" in package.namelist() and "xl/workbook.xml" in package.namelist(), "xlsx")
    again = apply_receipts(book, [receipt()], tmp / "receipts", tmp / "backups", name_prefix="TEST - ", stamp="20261024-120100")
    check(again["skipped"] == ["Costco|2026-10-03|12.34"], again)
    check(again["backup"] == "", "a duplicate must not replace the workbook")
    check(grid_amount(book, "Oct", "O6") == 12.34, "amount was added twice")
    totals = report_totals(book, date(2026, 10, 1), date(2026, 10, 7))
    check(totals.get("Meals") == 12.34, totals)
    check("Meals" in advice(totals), advice(totals))
    check(advice({}) == "", "empty report stays quiet")


def test_split_fills_grid(tmp: Path) -> None:
    book = new_book(tmp)
    item = receipt(
        total=14.34,
        items=[
            {"name": "milk", "price": 10.00, "category": "Meals"},
            {"name": "soap", "price": 4.34, "category": "Household Goods"},
        ],
    )
    apply_receipts(book, [item], tmp / "receipts", tmp / "backups", name_prefix="TEST - ", stamp="20261024-120200")
    check(grid_amount(book, "Oct", "O6") == 10.00, grid_amount(book, "Oct", "O6"))
    check(grid_amount(book, "Oct", "O22") == 4.34, grid_amount(book, "Oct", "O22"))
    check(transactions_empty(book), "split receipt is not a statement line")


def test_suggestion_remembers_no(tmp: Path) -> None:
    book = new_book(tmp)
    today = date(2026, 10, 24)
    check(suggest_statement("off", "not yet", 1, date(2026, 10, 23), today) == "", "too soon")
    early = suggest_statement("off", "not yet", 1, date(2026, 10, 10), today)
    check("GitHub" in early and "won't ask again" in early, early)
    check("own drive" in early, early)
    check(suggest_statement("off", "declined", 40, date(2026, 1, 1), today) == "", "a no sticks")
    check(suggest_statement("off", "asked", 40, date(2026, 1, 1), today) == "", "already asked")
    check(suggest_statement("on", "not yet", 40, date(2026, 1, 1), today) == "", "features already on")
    # Totals are far enough apart that rule 3 does not treat them as one purchase.
    spaced = [10, 30, 55, 90, 140, 210, 310, 450, 650, 950]
    for index, total in enumerate(spaced):
        result = apply_receipts(
            book,
            [receipt(when=f"2026-10-03 12:{index:02d}", total=total, body=f"Costco\nTotal {total:.2f}\n")],
            tmp / "receipts",
            tmp / "backups",
            name_prefix="TEST - ",
            stamp=f"20261024-13{index:04d}",
        )
        check(result["questions"] == [], result)
        check(result["wrote"], result)
    stats = receipt_stats(book)
    check(stats["count"] == 10, stats)
    started = date.fromisoformat(str(load_workbook(book)["Annual Summary"]["C52"].value)[:10])
    # The day clock uses Book started, not the Oct 3 receipt date.
    check(suggest_statement("off", "not yet", 1, started, today) == "", "a new book does not suggest from an old receipt date")
    text = suggest_statement("off", "not yet", stats["count"], started, today)
    check(text.startswith("You've been logging receipts"), text)
    before = snapshot_numbers(book)
    set_offer(book, "declined")
    check(snapshot_numbers(book) == before, "remembering a no changed a number")
    check(load_workbook(book)["Annual Summary"]["C48"].value == "declined", "offer")
    apply_receipts(
        book,
        [receipt(total=99.99, body="Costco\nTotal 99.99\n")],
        tmp / "receipts",
        tmp / "backups",
        name_prefix="TEST - ",
        stamp="20261024-140000",
    )
    later = receipt_stats(book)
    check(suggest_statement("off", "declined", later["count"], later["first"], today) == "", "must not nag")
    start, end, period = week_period(date(2026, 10, 4))
    check(period == "2026-10-04", period)
    check(start == date(2026, 9, 28) and end == date(2026, 10, 4), (start, end))
    record_report_sent(book, "weekly", period)
    check(report_already_sent(book, "weekly", period), "weekly report would be sent twice")


def test_queue_offline_lock_backup(tmp: Path) -> None:
    computer_root = tmp / "computer"
    book = new_book(computer_root)
    staging = tmp / "bot-staging"
    queue_path = tmp / "TEST - queue.json"
    queue = ReceiptQueue(queue_path)
    first = receipt()
    second_same = receipt(body="Costco\nTotal 12.34\nsecond photo\n")
    check(queue.enqueue(first)["result"] == "queued", "enqueue")
    check(queue.enqueue(second_same)["result"] == "duplicate", "same key queued twice")
    check(len(queue.items) == 1, queue.items)
    offline = LocalComputer(computer_root, online=False)
    waiting = drain(queue, offline, book.name, staging, name_prefix="TEST - ", stamp="20261024-150000")
    check(waiting["waiting"] == ["Costco|2026-10-03|12.34"], waiting)
    check(offline.approvals == 0, offline.approvals)
    check(grid_amount(book, "Oct", "O6") in (None, ""), "offline drain wrote the book")
    check(queue.items[0]["status"] == "queued", queue.items[0])
    online = LocalComputer(computer_root, online=True)
    done = drain(queue, online, book.name, staging, name_prefix="TEST - ", stamp="20261024-150000")
    check(done["wrote"] == ["Costco|2026-10-03|12.34"], done)
    check(online.approvals == 2, online.approvals)
    check(grid_amount(book, "Oct", "O6") == 12.34, grid_amount(book, "Oct", "O6"))
    backup = computer_root / "backups" / "TEST - Family Book - 2026 - 20261024-150000.xlsx"
    check(backup.exists(), "missing backup")
    check(load_workbook(backup)["Oct"]["O6"].value in (None, ""), "backup was taken after the write")
    check(load_workbook(book)["Oct"]["O6"].value == 12.34, "live book")
    quiet = drain(queue, online, book.name, staging, name_prefix="TEST - ", stamp="20261024-150100")
    check(quiet["wrote"] == [], quiet)
    check(grid_amount(book, "Oct", "O6") == 12.34, "second drain duplicated the amount")

    other = receipt(store="Shell", when="2026-10-04 09:15", total=40.00, category="Fuel", body="Shell\nTotal 40.00\n")
    queue.enqueue(other)
    lock = Path(str(book) + ".lock")
    lock.write_text("open in another app", encoding="utf-8")
    check(is_locked(book), "lock marker")
    held = drain(queue, online, book.name, staging, name_prefix="TEST - ", stamp="20261024-160000")
    check(held["locked"] is True, held)
    check(grid_amount(book, "Oct", "O6") == 12.34, "locked write changed the first receipt")
    # Day 4 Fuel is Q29. K29 is day 1 and is the wrong cell for this receipt.
    check(grid_amount(book, "Oct", "Q29") in (None, ""), grid_amount(book, "Oct", "Q29"))
    check(queue.items[-1]["status"] == "retry", queue.items)
    check(len(queue.items) == 2, "lock created a second queue row")
    lock.unlink()
    released = drain(queue, online, book.name, staging, name_prefix="TEST - ", stamp="20261024-160100")
    check(released["wrote"] == ["Shell|2026-10-04|40.00"], released)
    check(grid_amount(book, "Oct", "O6") == 12.34, "retry duplicated Costco")
    # Day 4 amount column is Q (17). Fuel is row 29.
    check(grid_amount(book, "Oct", "Q29") == 40.00, grid_amount(book, "Oct", "Q29"))
    released_again = drain(queue, online, book.name, staging, name_prefix="TEST - ", stamp="20261024-160200")
    check(released_again["wrote"] == [], released_again)
    check(grid_amount(book, "Oct", "Q29") == 40.00, "retry wrote fuel twice")


def test_batch_is_two_approvals(tmp: Path) -> None:
    computer_root = tmp / "computer"
    book = new_book(computer_root)
    queue = ReceiptQueue(tmp / "TEST - batch-queue.json")
    # 5, 20, and 40 are far enough apart that they are three purchases, not one near match.
    for total in (5, 20, 40):
        queue.enqueue(receipt(total=total, body=f"Costco\nTotal {total:.2f}\n"))
    computer = LocalComputer(computer_root, online=True)
    result = drain(queue, computer, book.name, tmp / "staging", name_prefix="TEST - ", stamp="20261024-170000")
    check(len(result["wrote"]) == 3, result)
    check(result["questions"] == [], result)
    check(computer.approvals == 2, computer.approvals)
    wb = load_workbook(book)
    amounts = [wb["Oct"].cell(6 + index, 15).value for index in range(3)]
    check(amounts == [5, 20, 40], amounts)


def test_settings_migration_keeps_amounts(tmp: Path) -> None:
    folder = tmp / "old"
    folder.mkdir(parents=True)
    book = new_book(folder)
    wb = load_workbook(book)
    wb["Oct"]["O6"] = 12.34
    wb["Oct"]["C19"] = 80
    wb["Annual Summary"]["C46"] = "v2.0"
    wb["Annual Summary"]["C47"] = None
    wb["Annual Summary"]["B47"] = None
    wb["Annual Summary"]["C48"] = None
    wb["Annual Summary"]["B2"] = "Family Book v2.0 by SkybatanX\nKeep this line."
    wb.save(book)
    before = snapshot_numbers(book)
    result = migrate_v21_settings(book)
    check(result["changed"] is True, result)
    after = snapshot_numbers(book)
    check(before == after, "migration changed a number or a formula")
    wb = load_workbook(book)
    check(wb["Annual Summary"]["C46"].value == "v2.1", wb["Annual Summary"]["C46"].value)
    check(wb["Annual Summary"]["C47"].value == "off", wb["Annual Summary"]["C47"].value)
    check(wb["Annual Summary"]["C48"].value == "not yet", wb["Annual Summary"]["C48"].value)
    check(str(wb["Annual Summary"]["B2"].value).startswith("Family Book v2.1"), wb["Annual Summary"]["B2"].value)
    check("Keep this line." in str(wb["Annual Summary"]["B2"].value), "banner body was dropped")
    again = migrate_v21_settings(book)
    check(again["changed"] is False, again)
    check(snapshot_numbers(book) == before, "rerun changed a number")

    wb = load_workbook(book)
    wb["Annual Summary"]["C46"] = "v2.0"
    wb["Annual Summary"]["C47"] = None
    wb["Annual Summary"]["C48"] = "declined"
    wb["Jan"]["B49"] = "CIBC Visa-1001"
    wb["Jan"]["G49"] = 10
    wb.save(book)
    kept = snapshot_numbers(book)
    migrate_v21_settings(book)
    check(snapshot_numbers(book) == kept, "statement book migration changed an amount")
    wb = load_workbook(book)
    check(wb["Annual Summary"]["C47"].value == "on", "a book that already has a statement stays on")
    check(wb["Annual Summary"]["C48"].value == "declined", "a no must survive migration")
    check(wb["Jan"]["G49"].value == 10, wb["Jan"]["G49"].value)


def test_near_match_and_rule2(tmp: Path) -> None:
    book = new_book(tmp)
    first = apply_receipts(book, [receipt()], tmp / "receipts", tmp / "backups", name_prefix="TEST - ", stamp="20261024-181000")
    check(first["wrote"] == ["Costco|2026-10-03|12.34"], first)
    later = receipt(when="2026-10-04 12:00", total=12.34, body="Costco\n2026-10-04\nTotal 12.34\n")
    attached = apply_receipts(book, [later], tmp / "receipts", tmp / "backups", name_prefix="TEST - ", stamp="20261024-181100")
    check(attached["attached"] == ["Costco|2026-10-04|12.34"], attached)
    check(attached["wrote"] == [], attached)
    check(grid_amount(book, "Oct", "O6") == 12.34, "rule 2 added a second amount on day 3")
    check(grid_amount(book, "Oct", "Q6") in (None, ""), grid_amount(book, "Oct", "Q6"))
    split = receipt(
        total=14.34,
        body="Costco\nTotal 14.34\nmilk 10.00\nsoap 4.34\n",
        items=[
            {"name": "milk", "price": 10.00, "category": "Meals"},
            {"name": "soap", "price": 4.34, "category": "Household Goods"},
        ],
    )
    asked = apply_receipts(book, [split], tmp / "receipts", tmp / "backups", name_prefix="TEST - ", stamp="20261024-181200")
    check(asked["wrote"] == [], asked)
    check(asked["questions"], asked)
    check("Same purchase as Costco $12.34 on Oct 3?" in asked["questions"][0]["text"], asked["questions"])
    check("Reply no if it's new." in asked["questions"][0]["text"], asked["questions"])
    check(grid_amount(book, "Oct", "O6") == 12.34, "near match changed the first amount")
    check(grid_amount(book, "Oct", "O7") in (None, ""), "near match wrote the next meals cell")
    check(grid_amount(book, "Oct", "O22") in (None, ""), "near match wrote household")
    log = load_workbook(book)["Log"]
    statuses = [log.cell(row, 14).value for row in range(2, 6)]
    check("question" in statuses, statuses)


def test_two_candidates_ask_which_date(tmp: Path) -> None:
    # Different receipt numbers make Oct 3 and Oct 11 two purchases. The Oct 10
    # slip has no number, so both are candidates and the bot must ask.
    fresh = new_book(tmp / "numbered")
    apply_receipts(
        fresh,
        [receipt(when="2026-10-03 09:00", total=57.29, number="100", body="Costco\n#100\nTotal 57.29\n")],
        tmp / "numbered-receipts",
        tmp / "numbered-backups",
        name_prefix="TEST - ",
        stamp="20261024-182200",
    )
    apply_receipts(
        fresh,
        [receipt(when="2026-10-11 09:00", total=57.29, number="200", body="Costco\n#200\nTotal 57.29\n")],
        tmp / "numbered-receipts",
        tmp / "numbered-backups",
        name_prefix="TEST - ",
        stamp="20261024-182300",
    )
    asked = apply_receipts(
        fresh,
        [receipt(when="2026-10-10 09:00", total=57.29, body="Costco\n2026-10-10\nTotal 57.29\n")],
        tmp / "numbered-receipts",
        tmp / "numbered-backups",
        name_prefix="TEST - ",
        stamp="20261024-182400",
    )
    check(len(asked["questions"]) == 1, asked)
    text = asked["questions"][0]["text"]
    check("Which date does it match:" in text and "Reply no if it's new." in text, text)
    check("Costco $57.29 on Oct 11" in text and "Costco $57.29 on Oct 3" in text, text)
    check(asked["wrote"] == [], asked)
    bare = answer_question(
        fresh,
        "Costco|2026-10-10|57.29",
        "yes",
        tmp / "numbered-receipts",
        tmp / "numbered-backups",
        name_prefix="TEST - ",
        stamp="20261024-182500",
    )
    check(bare["action"] == "ask", bare)
    check(grid_amount(fresh, "Oct", "AC6") in (None, ""), "a bare yes wrote day 10")
    named = answer_question(
        fresh,
        "Costco|2026-10-10|57.29",
        "Oct 11",
        tmp / "numbered-receipts",
        tmp / "numbered-backups",
        name_prefix="TEST - ",
        stamp="20261024-182600",
    )
    check(named["action"] == "attached", named)
    # The earlier of Oct 10 and Oct 11 keeps the same 57.29. It is not added again.
    check(grid_amount(fresh, "Oct", "AC6") == 57.29, grid_amount(fresh, "Oct", "AC6"))
    check(grid_amount(fresh, "Oct", "AE6") in (None, ""), grid_amount(fresh, "Oct", "AE6"))
    check(grid_amount(fresh, "Oct", "O6") == 57.29, grid_amount(fresh, "Oct", "O6"))


def test_backup_stamp_is_not_a_lock(tmp: Path) -> None:
    computer_root = tmp / "computer"
    book = new_book(computer_root)
    queue = ReceiptQueue(tmp / "TEST - stamp-queue.json")
    queue.enqueue(receipt())
    computer = LocalComputer(computer_root, online=True)
    first = drain(queue, computer, book.name, tmp / "staging-a", name_prefix="TEST - ", stamp="20261024-120000")
    check(first["locked"] is False, first)
    check(grid_amount(book, "Oct", "O6") == 12.34, "first drain")
    queue.enqueue(receipt(store="Shell", when="2026-10-04 09:15", total=40.00, category="Fuel", body="Shell\nTotal 40.00\n"))
    second = drain(queue, computer, book.name, tmp / "staging-b", name_prefix="TEST - ", stamp="20261024-120000")
    check(second["locked"] is False, second)
    check(second["wrote"] == ["Shell|2026-10-04|40.00"], second)
    check(grid_amount(book, "Oct", "Q29") == 40.00, grid_amount(book, "Oct", "Q29"))
    check(grid_amount(book, "Oct", "O6") == 12.34, "stamp collision duplicated Costco")
    backups = list((computer_root / "backups").glob("*.xlsx"))
    check(len(backups) == 2, [path.name for path in backups])
    check(all(queue_item["status"] == "done" for queue_item in queue.items), queue.items)


def _split_receipt() -> dict:
    return receipt(
        total=14.34,
        body="Costco\nTotal 14.34\nmilk 10.00\nsoap 4.34\n",
        items=[
            {"name": "milk", "price": 10.00, "category": "Meals"},
            {"name": "soap", "price": 4.34, "category": "Household Goods"},
        ],
    )


def test_answer_no_yes_and_rerun(tmp: Path) -> None:
    key = "Costco|2026-10-03|14.34"
    computer_root = tmp / "no-computer"
    book = new_book(computer_root)
    queue = ReceiptQueue(tmp / "TEST - answer-no.json")
    computer = LocalComputer(computer_root, online=True)
    queue.enqueue(receipt())
    drain(queue, computer, book.name, tmp / "stage-no-1", name_prefix="TEST - ", stamp="20261029-100000")
    queue.enqueue(_split_receipt())
    asked = drain(queue, computer, book.name, tmp / "stage-no-2", name_prefix="TEST - ", stamp="20261029-100100")
    check(asked["questions"], asked)
    text = asked["questions"][0]["text"]
    check("Same purchase as Costco $12.34 on Oct 3?" in text, text)
    check("Reply no if it's new." in text, text)
    check(grid_amount(book, "Oct", "O7") in (None, ""), "the question wrote before the answer")
    answered = resolve_question(
        queue, computer, book.name, tmp / "stage-no-3", key, "no, it's new",
        name_prefix="TEST - ", stamp="20261029-100200",
    )
    check(answered["action"] == "wrote", answered)
    check(answered["locked"] is False, answered)
    check(grid_amount(book, "Oct", "O6") == 12.34, grid_amount(book, "Oct", "O6"))
    check(grid_amount(book, "Oct", "O7") == 10.00, grid_amount(book, "Oct", "O7"))
    check(grid_amount(book, "Oct", "O22") == 4.34, grid_amount(book, "Oct", "O22"))
    statuses = [load_workbook(book)["Log"].cell(row, 14).value for row in range(2, 6)]
    check("question" not in statuses and "cash" in statuses, statuses)
    check(all(item["status"] == "done" for item in queue.items), queue.items)
    again = apply_receipts(
        book, [_split_receipt()], computer_root / "receipts", computer_root / "backups",
        name_prefix="TEST - ", stamp="20261029-100300",
    )
    check(again["wrote"] == [], again)
    check(again["skipped"] == [key], again)
    check(grid_amount(book, "Oct", "O7") == 10.00, "rerun wrote meals again")
    check(grid_amount(book, "Oct", "O22") == 4.34, "rerun wrote household again")
    changed_after_no = _split_receipt()
    changed_after_no["body"] = "Costco\nTotal 14.34\nmilk 10.00\nsoap 4.34\nrevised after no\n"
    attached_file = apply_receipts(
        book, [changed_after_no], computer_root / "receipts", computer_root / "backups",
        name_prefix="TEST - ", stamp="20261029-100350",
    )
    check(attached_file["attached"] == [key] and attached_file["wrote"] == [], attached_file)
    check(grid_amount(book, "Oct", "O6") == 12.34, "changed body after a no added 12.34")
    check(grid_amount(book, "Oct", "O7") == 10.00, "changed body after a no added meals")
    check(grid_amount(book, "Oct", "O22") == 4.34, "changed body after a no added household")
    repeat = answer_question(
        book, key, "no", computer_root / "receipts", computer_root / "backups",
        name_prefix="TEST - ", stamp="20261029-100400",
    )
    check(repeat["action"] == "skip", repeat)
    check(grid_amount(book, "Oct", "O7") == 10.00, "a second no wrote again")

    yes_root = tmp / "yes-computer"
    yes_book = new_book(yes_root)
    apply_receipts(yes_book, [receipt()], yes_root / "receipts", yes_root / "backups", name_prefix="TEST - ", stamp="20261029-110000")
    apply_receipts(yes_book, [_split_receipt()], yes_root / "receipts", yes_root / "backups", name_prefix="TEST - ", stamp="20261029-110100")
    yes = answer_question(
        yes_book, key, "yes", yes_root / "receipts", yes_root / "backups",
        name_prefix="TEST - ", stamp="20261029-110200",
    )
    check(yes["action"] == "attached", yes)
    check(grid_amount(yes_book, "Oct", "O6") == 12.34, "yes changed the first amount")
    check(grid_amount(yes_book, "Oct", "O7") in (None, ""), grid_amount(yes_book, "Oct", "O7"))
    check(grid_amount(yes_book, "Oct", "O22") in (None, ""), grid_amount(yes_book, "Oct", "O22"))
    yes_status = [load_workbook(yes_book)["Log"].cell(row, 14).value for row in range(2, 6)]
    check("question" not in yes_status and "superseded" in yes_status, yes_status)
    rerun = apply_receipts(
        yes_book, [_split_receipt()], yes_root / "receipts", yes_root / "backups",
        name_prefix="TEST - ", stamp="20261029-110300",
    )
    check(rerun["wrote"] == [], rerun)
    changed = _split_receipt()
    changed["body"] = "Costco\nTotal 14.34\nmilk 10.00\nsoap 4.34\nrevised\n"
    revised = apply_receipts(
        yes_book, [changed], yes_root / "receipts", yes_root / "backups",
        name_prefix="TEST - ", stamp="20261029-110400",
    )
    check(revised["wrote"] == [] and revised["skipped"] == [key], revised)
    check(grid_amount(yes_book, "Oct", "O6") == 12.34, "revised body added an amount")
    check(grid_amount(yes_book, "Oct", "O7") in (None, ""), "revised body filled O7")

    date_root = tmp / "date-computer"
    date_book = new_book(date_root)
    apply_receipts(date_book, [receipt()], date_root / "receipts", date_root / "backups", name_prefix="TEST - ", stamp="20261029-120000")
    apply_receipts(date_book, [_split_receipt()], date_root / "receipts", date_root / "backups", name_prefix="TEST - ", stamp="20261029-120100")
    dated = answer_question(
        date_book, key, "Oct 3", date_root / "receipts", date_root / "backups",
        name_prefix="TEST - ", stamp="20261029-120200",
    )
    check(dated["action"] == "attached", dated)
    check(grid_amount(date_book, "Oct", "O6") == 12.34, grid_amount(date_book, "Oct", "O6"))
    check(grid_amount(date_book, "Oct", "O7") in (None, ""), "a date wrote a new meals cell")
    check(grid_amount(date_book, "Oct", "O22") in (None, ""), "a date wrote household")
    again_date = apply_receipts(
        date_book, [_split_receipt()], date_root / "receipts", date_root / "backups",
        name_prefix="TEST - ", stamp="20261029-120300",
    )
    check(again_date["wrote"] == [], again_date)
    check(grid_amount(date_book, "Oct", "O6") == 12.34, "date rerun duplicated 12.34")


def _copy_book(source: Path, folder: Path) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / source.name
    shutil.copy(source, dest)
    return dest


def test_reply_phrases(tmp: Path) -> None:
    """Dates win over a leading 'no'. A whole reply is yes or no. Anything else asks again."""
    october = new_book(tmp / "october")
    apply_receipts(
        october, [receipt()], tmp / "october-receipts", tmp / "october-backups",
        name_prefix="TEST - ", stamp="20261029-140000",
    )
    asked = apply_receipts(
        october, [_split_receipt()], tmp / "october-receipts", tmp / "october-backups",
        name_prefix="TEST - ", stamp="20261029-140100",
    )
    check(asked["questions"], asked)
    october_key = "Costco|2026-10-03|14.34"
    negatives = ("no", "nope", "nah", "no, it's new", "new", "different")
    affirmatives = ("yes, same purchase", "yes it matches", "same", "that one", "correct")
    for index, phrase in enumerate(negatives):
        book = _copy_book(october, tmp / f"no-{index}")
        result = answer_question(
            book, october_key, phrase, tmp / "october-receipts", tmp / f"no-backups-{index}",
            name_prefix="TEST - ", stamp=f"20261029-141{index}00",
        )
        check(result["action"] == "wrote", (phrase, result))
        check(grid_amount(book, "Oct", "O6") == 12.34, phrase)
        check(grid_amount(book, "Oct", "O7") == 10.00, (phrase, grid_amount(book, "Oct", "O7")))
        check(grid_amount(book, "Oct", "O22") == 4.34, (phrase, grid_amount(book, "Oct", "O22")))
        statuses = [load_workbook(book)["Log"].cell(row, 14).value for row in range(2, 6)]
        check("question" not in statuses, (phrase, statuses))
    for index, phrase in enumerate(affirmatives):
        book = _copy_book(october, tmp / f"yes-{index}")
        result = answer_question(
            book, october_key, phrase, tmp / "october-receipts", tmp / f"yes-backups-{index}",
            name_prefix="TEST - ", stamp=f"20261029-142{index}00",
        )
        check(result["action"] == "attached", (phrase, result))
        check(grid_amount(book, "Oct", "O6") == 12.34, phrase)
        check(grid_amount(book, "Oct", "O7") in (None, ""), (phrase, grid_amount(book, "Oct", "O7")))
        check(grid_amount(book, "Oct", "O22") in (None, ""), (phrase, grid_amount(book, "Oct", "O22")))
        statuses = [load_workbook(book)["Log"].cell(row, 14).value for row in range(2, 6)]
        check("question" not in statuses and "superseded" in statuses, (phrase, statuses))
    ambiguous = _copy_book(october, tmp / "ambiguous")
    note = answer_question(
        ambiguous, october_key, "note: same purchase", tmp / "october-receipts", tmp / "note-backups",
        name_prefix="TEST - ", stamp="20261029-143000",
    )
    check(note["action"] == "ask" and note["backup"] == "", note)
    check(grid_amount(ambiguous, "Oct", "O6") == 12.34, "an ambiguous reply wrote")
    check(grid_amount(ambiguous, "Oct", "O7") in (None, ""), grid_amount(ambiguous, "Oct", "O7"))
    check(grid_amount(ambiguous, "Oct", "O22") in (None, ""), grid_amount(ambiguous, "Oct", "O22"))
    note_status = [load_workbook(ambiguous)["Log"].cell(row, 14).value for row in range(2, 6)]
    check("question" in note_status, note_status)

    november = new_book(tmp / "november")
    apply_receipts(
        november,
        [receipt(when="2026-11-03 12:00", total=12.34, body="Costco\n2026-11-03\nTotal 12.34\n")],
        tmp / "november-receipts", tmp / "november-backups",
        name_prefix="TEST - ", stamp="20261029-150000",
    )
    november_near = receipt(when="2026-11-03 12:00", total=13.00, body="Costco\n2026-11-03\nTotal 13.00\n")
    november_asked = apply_receipts(
        november, [november_near], tmp / "november-receipts", tmp / "november-backups",
        name_prefix="TEST - ", stamp="20261029-150100",
    )
    check(november_asked["questions"] and november_asked["wrote"] == [], november_asked)
    november_key = "Costco|2026-11-03|13.00"
    for index, phrase in enumerate(("Nov 3", "November 3", "2026-11-03")):
        book = _copy_book(november, tmp / f"date-{index}")
        result = answer_question(
            book, november_key, phrase, tmp / "november-receipts", tmp / f"date-backups-{index}",
            name_prefix="TEST - ", stamp=f"20261029-151{index}00",
        )
        check(result["action"] == "attached", (phrase, result))
        check(grid_amount(book, "Nov", "O6") == 12.34, (phrase, grid_amount(book, "Nov", "O6")))
        check(grid_amount(book, "Nov", "O7") in (None, ""), (phrase, grid_amount(book, "Nov", "O7")))
        statuses = [load_workbook(book)["Log"].cell(row, 14).value for row in range(2, 6)]
        check("question" not in statuses and "superseded" in statuses, (phrase, statuses))


def test_near_total_one_day_apart(tmp: Path) -> None:
    book = new_book(tmp)
    apply_receipts(
        book,
        [receipt(when="2026-10-06 09:00", total=41.00, body="Costco\n2026-10-06\nTotal 41.00\n")],
        tmp / "receipts",
        tmp / "backups",
        name_prefix="TEST - ",
        stamp="20261029-130000",
    )
    asked = apply_receipts(
        book,
        [receipt(when="2026-10-07 09:00", total=42.00, body="Costco\n2026-10-07\nTotal 42.00\n")],
        tmp / "receipts",
        tmp / "backups",
        name_prefix="TEST - ",
        stamp="20261029-130100",
    )
    check(asked["wrote"] == [] and asked["attached"] == [], asked)
    check(len(asked["questions"]) == 1, asked)
    check("Costco $41.00 on Oct 6" in asked["questions"][0]["text"], asked["questions"])
    check(grid_amount(book, "Oct", "U6") == 41.00, grid_amount(book, "Oct", "U6"))
    check(grid_amount(book, "Oct", "W6") in (None, ""), grid_amount(book, "Oct", "W6"))


def test_first_statement_turns_features_on(tmp: Path) -> None:
    book = new_book(tmp)
    apply_receipts(book, [receipt()], tmp / "receipts", tmp / "backups", name_prefix="TEST - ", stamp="20261024-180000")
    before = snapshot_numbers(book)
    check(routine_gate(book)["may_remind"] is False, "quiet before any statement")
    set_offer(book, "declined")
    mark_statement_added(book)
    check(load_workbook(book)["Annual Summary"]["C48"].value == "accepted", "an import the user chose flips a no to accepted")
    check(snapshot_numbers(book) == before, "enabling statements changed an amount")
    gate = routine_gate(book)
    check(gate["features"] == "on" and gate["may_remind"] and gate["may_nudge"] and gate["may_bank_income"], gate)
    check(grid_amount(book, "Oct", "O6") == 12.34, "receipt amount moved")


def test_v2_statement_flow_still_posts_once() -> None:
    text = (ROOT / "tests" / "v2.0" / "fixtures" / "mid-month.csv").read_text(encoding="utf-8")
    book = V2.Book()
    statement = V2.parse_statement(text, "csv")
    book.import_statement(statement, "mid-month.csv")
    book.import_statement(statement, "mid-month.csv")
    check(book.cell("Aug", "BM29") == 71.04, book.grid)
    check(len(book.txns) == 1, book.txns)


def test_template_version() -> None:
    wb = load_workbook(ROOT / "template" / "family-book-template.xlsx")
    summary = wb["Annual Summary"]
    check(summary["C46"].value == "v2.1", summary["C46"].value)
    check(summary["B47"].value == "Statement features", summary["B47"].value)
    check(summary["C47"].value == "off", summary["C47"].value)
    check(summary["C48"].value == "not yet", summary["C48"].value)
    check("v2.1" in str(summary["B2"].value), summary["B2"].value)
    check("Excel does not have to be installed" in str(summary["B2"].value), summary["B2"].value)


def main() -> None:
    test_keys_match_cloud()
    test_v2_statement_flow_still_posts_once()
    with tempfile.TemporaryDirectory(prefix="TEST - family-book-") as directory:
        root = Path(directory)
        test_receipts_only_and_no_excel(root / "plain")
        test_split_fills_grid(root / "split")
        test_suggestion_remembers_no(root / "suggest")
        test_queue_offline_lock_backup(root / "queue")
        test_batch_is_two_approvals(root / "batch")
        test_settings_migration_keeps_amounts(root / "migrate")
        test_near_match_and_rule2(root / "near")
        test_two_candidates_ask_which_date(root / "which")
        test_answer_no_yes_and_rerun(root / "answer")
        test_reply_phrases(root / "phrases")
        test_near_total_one_day_apart(root / "near-day")
        test_backup_stamp_is_not_a_lock(root / "stamp")
        test_first_statement_turns_features_on(root / "enable")
    test_template_version()
    # Reuse the v2.0 checker without regenerating its fixtures.
    import importlib.util
    path = ROOT / "tests" / "v2.0" / "generate_and_check.py"
    spec = importlib.util.spec_from_file_location("generate_and_check", path)
    checker = importlib.util.module_from_spec(spec)
    sys.modules["generate_and_check"] = checker
    spec.loader.exec_module(checker)
    checker.test_parsers()
    checker.test_book()
    checker.test_round1()
    checker.test_migration_and_crash()
    checker.test_template()
    print("v2.1 offline checks passed")


if __name__ == "__main__":
    try:
        main()
    except Fail as exc:
        print("FAIL:", exc)
        sys.exit(1)
