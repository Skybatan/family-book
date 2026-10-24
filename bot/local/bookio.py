#!/usr/bin/env python3
"""Read and write a Family Book .xlsx without Microsoft Excel.

The bot runs this on its own computer after copying the workbook off the
user's computer. Nothing here starts the Excel application. A locked file
is left untouched. A receipt key that is already in the book is not written
again. A backup is taken before a workbook is replaced.

Receipt key, categories, and the month-sheet layout match the cloud books.
"""

from __future__ import annotations

import hashlib
import re
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# Daily Record rows. These match the family-book skill and the template.
CATEGORIES = {
    "Meals": (6, 16),
    "Snacks & Groceries": (17, 19),
    "Dining Out": (20, 20),
    "Household Goods": (22, 27),
    "Education & Enrichment": (28, 28),
    "Fuel": (29, 29),
    "Clothing": (30, 31),
    "Medical": (32, 32),
    "Hair & Cosmetics": (33, 33),
    "Transportation": (34, 34),
    "Entertainment": (35, 35),
    "Gifts & Social": (36, 36),
    "Other": (37, 40),
    "Temporary Income": (43, 43),
}

OFFER_VALUES = ("not yet", "asked", "declined", "accepted")
FEATURE_VALUES = ("off", "on")

SUGGESTION_TEXT = (
    "You've been logging receipts for a while. If you want, you can also add a "
    "card or bank statement. That catches purchases you didn't photograph, "
    "reminds you before a card is due, and can fill in income and fixed bills. "
    "The original file stays in your own drive. Nothing is uploaded to GitHub. "
    "Say yes and send a statement, or say no and I won't ask again."
)

BANNER_V21 = (
    "Family Book v2.1 by SkybatanX: a yearly family income & expense book run by your personal Grok Bot.\n"
    "The bot logs spending from receipts. Card and bank statements are optional, and the book works with none.\n"
    "Each month keeps the number grid, then a Transactions section and a Receipts section. The same purchase is never added twice.\n"
    "The book can live in Google Sheets, Excel Online, or an Excel file on your computer. Excel does not have to be installed.\n"
    "Follow @SkyXFunNext on X for more templates and fun!"
)


def short_store(name: str) -> str:
    """Short merchant name used in the Store column, the file name, and the receipt key."""
    text = (name or "").strip()
    folded = text.lower().replace("&", " and ")
    folded = re.sub(r"#\s*\d+", " ", folded)
    folded = re.sub(r"\b(no\.?|store)\s*\d+\b", " ", folded)
    folded = re.sub(r"\.(com|ca|co\.uk|net|org)\b", " ", folded)
    folded = re.sub(r"\b(wholesale|online|the)\b", " ", folded)
    folded = re.sub(r"[^a-z0-9\s]", " ", folded)
    folded = re.sub(r"\s+", " ", folded).strip()
    brands = (
        ("costco", "Costco"),
        ("home depot", "Home Depot"),
        ("shell", "Shell"),
        ("walmart", "Walmart"),
        ("canadian tire", "Canadian Tire"),
        ("sport chek", "Sport Chek"),
        ("tim hortons", "Tim Hortons"),
        ("t and t", "T&T"),
    )
    for needle, pretty in brands:
        if needle in folded:
            return pretty
    if not folded:
        return "Receipt"
    return " ".join(part.capitalize() for part in folded.split()[:3])


def normalize_number(number: str) -> str:
    text = (number or "").strip()
    if text.startswith("#"):
        text = text[1:].strip()
    return text.casefold()


def receipt_key(store: str, number: str, when, total) -> str:
    """Same key as the cloud books.

    A receipt number at that store is the key. With no number, the key is the
    short store name, the transaction date, and the total to the cent.
    """
    merchant = short_store(store)
    normalized = normalize_number(number)
    if normalized:
        return f"{merchant}|{normalized}"
    day = _as_date(when)
    day_text = day.isoformat() if day else ""
    return f"{merchant}|{day_text}|{float(total):.2f}"


def _as_date(when) -> date | None:
    if when is None or when == "":
        return None
    if isinstance(when, datetime):
        return when.date()
    if isinstance(when, date):
        return when
    text = str(when).strip()
    if not text:
        return None
    return date.fromisoformat(text[:10])


def _as_hhmm(when) -> str:
    if isinstance(when, datetime):
        return when.strftime("%H%M")
    text = str(when or "")
    match = re.search(r"(\d{2}):(\d{2})", text)
    if match:
        return match.group(1) + match.group(2)
    match = re.search(r"\b(\d{4})\b", text[10:])
    if match:
        return match.group(1)
    return ""


def file_name(when, store: str, total: float, ext: str, prefix: str = "") -> str:
    """YYYY-MM-DD HHMM - Store - TOTAL.ext, with an optional test prefix."""
    day = _as_date(when)
    day_text = day.isoformat() if day else "0000-00-00"
    hhmm = _as_hhmm(when)
    stamp = f"{day_text} {hhmm}" if hhmm else day_text
    suffix = ext if ext.startswith(".") else f".{ext}"
    total_text = f"{float(total):.2f}"
    base = f"{stamp} - {short_store(store)} - {total_text}{suffix}"
    safe = re.sub(r'[\\/:*?"<>|]', "", base)
    return f"{prefix}{safe}" if prefix else safe


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def is_locked(path: Path) -> bool:
    """True when another app has the workbook open, or a tester lock marker is present.

    Excel's owner file is ``~$<name>``. A ``.lock`` file beside the workbook is
    the stand-in a tester uses when Excel is not installed. A lock means the
    caller must not replace the file.
    """
    path = Path(path)
    if not path.exists():
        return False
    if path.with_name("~$" + path.name).exists():
        return True
    if Path(str(path) + ".lock").exists():
        return True
    try:
        import fcntl
        import os
        fd = os.open(path, os.O_RDWR)
    except OSError:
        return True
    try:
        try:
            fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            return True
        fcntl.flock(fd, fcntl.LOCK_UN)
        return False
    finally:
        os.close(fd)


def backup_workbook(path: Path, backups_dir: Path, stamp: str | None = None) -> Path:
    """Copy the workbook into backups before it is replaced. Never overwrite a backup."""
    path = Path(path)
    backups_dir = Path(backups_dir)
    backups_dir.mkdir(parents=True, exist_ok=True)
    when = stamp or datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = backups_dir / f"{path.stem} - {when}{path.suffix}"
    number = 2
    while dest.exists():
        dest = backups_dir / f"{path.stem} - {when} - {number}{path.suffix}"
        number += 1
    shutil.copy2(path, dest)
    if dest.stat().st_size <= 0:
        raise RuntimeError(f"backup is empty: {dest}")
    return dest


def _atomic_save(wb, path: Path) -> None:
    partial = Path(str(path) + ".partial")
    wb.save(partial)
    partial.replace(path)


def init_local_book(template: Path, folder: Path, year: int, workbook_name: str) -> Path:
    """Create a year folder, a receipts folder, and a year workbook. Excel is not used."""
    folder = Path(folder)
    receipts = folder / "receipts"
    backups = folder / "backups"
    receipts.mkdir(parents=True, exist_ok=True)
    backups.mkdir(parents=True, exist_ok=True)
    dest = folder / workbook_name
    if dest.exists():
        return dest
    shutil.copy2(template, dest)
    wb = load_workbook(dest)
    summary = wb["Annual Summary"]
    summary["F5"] = int(year)
    if summary["C46"].value in (None, ""):
        summary["C46"] = "v2.1"
    if summary["C47"].value not in FEATURE_VALUES:
        summary["B47"] = "Statement features"
        summary["C47"] = "off"
    if summary["C48"].value not in OFFER_VALUES:
        summary["B48"] = "Statement offer"
        summary["C48"] = "not yet"
    summary["B49"] = "Last weekly report"
    summary["B50"] = "Last monthly report"
    summary["B51"] = "Last yearly report"
    summary["B52"] = "Book started"
    if not summary["C52"].value:
        summary["C52"] = date.today().isoformat()
    _atomic_save(wb, dest)
    return dest


def _summary(wb):
    return wb["Annual Summary"]


def read_settings(wb) -> dict:
    summary = _summary(wb)
    features = summary["C47"].value if summary["C47"].value in FEATURE_VALUES else "off"
    offer = summary["C48"].value if summary["C48"].value in OFFER_VALUES else "not yet"
    return {
        "version": summary["C46"].value or "",
        "features": features,
        "offer": offer,
        "weekly": summary["C49"].value or "",
        "monthly": summary["C50"].value or "",
        "yearly": summary["C51"].value or "",
    }


def _sheet_for(when) -> tuple[str, date, int]:
    day = _as_date(when)
    if day is None:
        raise ValueError("receipt has no transaction date")
    return MONTHS[day.month - 1], day, day.day


def _day_columns(day_number: int) -> tuple[int, int]:
    item = 10 + 2 * (day_number - 1)
    return item, item + 1


def _items_of(receipt: dict) -> list[dict]:
    items = receipt.get("items") or []
    if items:
        return items
    category = receipt.get("category") or "Meals"
    if category not in CATEGORIES:
        raise ValueError(f"unknown category: {category}")
    return [{"name": short_store(receipt.get("store", "")), "price": float(receipt["total"]), "category": category}]


def _check_split(receipt: dict) -> list[dict]:
    items = _items_of(receipt)
    for item in items:
        if item.get("category") not in CATEGORIES:
            raise ValueError(f"unknown category: {item.get('category')}")
    total = round(sum(float(item["price"]) for item in items), 2)
    if abs(total - round(float(receipt["total"]), 2)) > 0.001:
        raise ValueError("item prices do not add up to the receipt total")
    return items


def _iter_receipt_rows(wb):
    for name in MONTHS:
        ws = wb[name]
        for row in range(152, 302):
            store = ws.cell(row, 3).value
            key = ws.cell(row, 79).value
            if store or key:
                yield ws, row


def _next_receipt_row(wb, sheet: str) -> int:
    ws = wb[sheet]
    for row in range(152, 302):
        if not ws.cell(row, 3).value and not ws.cell(row, 79).value:
            return row
    raise RuntimeError(f"no empty receipt row on {sheet}")


def _next_log_row(wb) -> int:
    log = wb["Log"]
    for row in range(2, 5000):
        if not log.cell(row, 3).value and not log.cell(row, 14).value:
            return row
    raise RuntimeError("no empty log row")


def _first_empty(ws, day_number: int, category: str, reserved: set[tuple]) -> tuple[int, int]:
    start, end = CATEGORIES[category]
    _item_col, amount_col = _day_columns(day_number)
    for row in range(start, end + 1):
        if (ws.title, row, amount_col) in reserved:
            continue
        value = ws.cell(row, amount_col).value
        if value is None or value == "":
            return row, amount_col
    raise RuntimeError(f"no empty {category} cell on day {day_number}")


def _placed_amount(ws, day_number: int, category: str, amount: float, store: str, reserved: set[tuple]):
    """A cell in this day's category block that already holds this amount and names the store."""
    start, end = CATEGORIES[category]
    item_col, amount_col = _day_columns(day_number)
    needle = short_store(store).casefold()
    for row in range(start, end + 1):
        if (ws.title, row, amount_col) in reserved:
            continue
        value = ws.cell(row, amount_col).value
        text = str(ws.cell(row, item_col).value or "")
        if isinstance(value, (int, float)) and abs(float(value) - amount) < 0.001 and needle in text.casefold():
            return row, amount_col
    return None


def _amount_cells(ws, day_number: int, items: list[dict], reserved: set[tuple] | None = None, store: str = "") -> tuple[list[tuple[str, int, int, float]], bool]:
    """Choose amount cells. Reuse a cell that already holds this amount, otherwise the first empty one.

    The bool is true when at least one cell is still empty and will be written.
    """
    reserved = reserved if reserved is not None else set()
    grouped: dict[str, float] = {}
    for item in items:
        grouped[item["category"]] = round(grouped.get(item["category"], 0) + float(item["price"]), 2)
    chosen = []
    fresh = False
    for category, amount in grouped.items():
        placed = _placed_amount(ws, day_number, category, amount, store, reserved) if store else None
        if placed:
            row, col = placed
        else:
            row, col = _first_empty(ws, day_number, category, reserved)
            fresh = True
        reserved.add((ws.title, row, col))
        chosen.append((category, row, col, amount))
    return chosen, fresh


def _where_text(sheet: str, cells: list[tuple[str, int, int, float]]) -> str:
    parts = []
    for _category, row, col, _amount in cells:
        parts.append(f"{sheet}!{get_column_letter(col)}{row}")
    return "+".join(parts)


def _grid_snapshot(wb) -> dict:
    """Numbers and formulas in the grid. Migration and a skipped receipt must leave these alone."""
    found = {}
    for name in MONTHS:
        ws = wb[name]
        for row in list(range(8, 12)) + list(range(18, 40)) + list(range(6, 44)):
            for col in range(1, 72):
                value = ws.cell(row, col).value
                if isinstance(value, (int, float)) or (isinstance(value, str) and value.startswith("=")):
                    found[(name, row, col)] = value
    return found


def _body_bytes(receipt: dict) -> bytes:
    body = receipt.get("body", b"")
    if isinstance(body, str):
        body = body.encode("utf-8")
    return body or b""


def _place_file(receipts_dir: Path, name: str, data: bytes) -> str:
    receipts_dir.mkdir(parents=True, exist_ok=True)
    digest = sha256_bytes(data) if data else ""
    candidate = receipts_dir / name
    if candidate.exists() and digest and sha256_bytes(candidate.read_bytes()) == digest:
        return candidate.name
    if candidate.exists():
        stem = candidate.stem
        suffix = candidate.suffix
        number = 2
        while candidate.exists():
            candidate = receipts_dir / f"{stem} - {number}{suffix}"
            number += 1
    if data:
        candidate.write_bytes(data)
    return candidate.name


def _find_log_row(wb, key: str):
    log = wb["Log"]
    for row in range(2, max(log.max_row, 2) + 1):
        merchant = log.cell(row, 3).value
        if not merchant:
            continue
        built = receipt_key(
            str(merchant),
            str(log.cell(row, 9).value or ""),
            log.cell(row, 2).value,
            log.cell(row, 4).value or 0,
        )
        if built == key:
            return row
    return None


def _find_receipt_row(wb, key: str):
    for ws, row in _iter_receipt_rows(wb):
        if ws.cell(row, 79).value == key:
            return ws, row
    return None, None


def merchant_match(left: str, right: str) -> str:
    """Return ``full``, ``partial``, or ``""`` using the same name rules as the cloud book."""

    def norm(value: str) -> str:
        text = (value or "").lower().replace("&", " and ")
        text = re.sub(r"#\s*\d+", " ", text)
        text = text.replace("*", " ")
        text = re.sub(r"\b(no\.?|store)\s*\d+\b", " ", text)
        text = re.sub(r"\.(com|ca|co\.uk|net|org)\b", " ", text)
        text = re.sub(r"\b(wholesale|online)\b", " ", text)
        text = re.sub(r"^\s*the\s+", " ", text)
        text = re.sub(r"[^a-z0-9\s]", " ", text)
        return re.sub(r"\s+", " ", text).strip()

    na, nb = norm(left), norm(right)
    if not na or not nb:
        return ""
    if na in nb or nb in na:
        return "full"
    words_a, words_b = na.split(), nb.split()
    shorter, longer = (words_a, words_b) if len(words_a) <= len(words_b) else (words_b, words_a)
    if shorter and all(word in longer for word in shorter):
        return "full"
    if any(len(word) >= 4 and word in words_b and word in words_a for word in words_a):
        return "partial"
    if abs(len(na) - len(nb)) <= 2 and _edits(na, nb) <= 2:
        return "partial"
    return ""


def _edits(left: str, right: str) -> int:
    if abs(len(left) - len(right)) > 2:
        return 99
    prev = list(range(len(right) + 1))
    for i, ca in enumerate(left, 1):
        cur = [i]
        for j, cb in enumerate(right, 1):
            cur.append(min(cur[-1] + 1, prev[j] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def _merchant_unclear(store: str) -> bool:
    text = (store or "").strip().casefold()
    return text in {"", "unknown", "store", "merchant", "receipt", "?"}


def _spoken_date(day: date) -> str:
    return f"{MONTHS[day.month - 1]} {day.day}"


def _money(amount) -> str:
    return f"{float(amount):.2f}"


def _log_rows(wb) -> list[dict]:
    log = wb["Log"]
    rows = []
    for row in range(2, max(log.max_row, 2) + 1):
        merchant = log.cell(row, 3).value
        status = log.cell(row, 14).value
        if not merchant and not status:
            continue
        total = log.cell(row, 4).value
        rows.append({
            "row": row,
            "source": str(log.cell(row, 1).value or ""),
            "date": _as_date(log.cell(row, 2).value),
            "merchant": str(merchant or ""),
            "total": float(total) if isinstance(total, (int, float)) else 0.0,
            "where": str(log.cell(row, 5).value or ""),
            "number": str(log.cell(row, 9).value or ""),
            "status": str(status or ""),
            "notes": str(log.cell(row, 8).value or ""),
            "hash": str(log.cell(row, 15).value or ""),
        })
    return rows


def _hash_seen(rows: list[dict], digest: str) -> bool:
    if not digest:
        return False
    return any(digest in row["hash"] for row in rows)


def _question_dates(row: dict) -> list[date]:
    found = []
    for text in (row["notes"], row["source"]):
        for match in re.findall(r"\d{4}-\d{2}-\d{2}", text):
            day = date.fromisoformat(match)
            if day not in found:
                found.append(day)
    if row["date"] and row["date"] not in found:
        found.append(row["date"])
    return found


def _near_limit(left: float, right: float) -> float:
    return max(5.0, 0.15 * max(abs(left), abs(right)))


def _gather_candidates(rows: list[dict], receipt: dict) -> list[dict]:
    """Every Log row that fits rule 1, 2, or 3. Question and superseded rows are not candidates."""
    store = receipt.get("store", "")
    number = normalize_number(receipt.get("number", ""))
    total = round(float(receipt["total"]), 2)
    day = _as_date(receipt.get("when"))
    found = []
    for row in rows:
        if row["status"] in {"superseded", "pending", "question", "refused"}:
            continue
        other_number = normalize_number(row["number"])
        if number and other_number and number != other_number:
            continue
        kind = merchant_match(row["merchant"], store)
        if not row["date"] or not day:
            continue
        gap = abs((row["date"] - day).days)
        same_total = abs(row["total"] - total) < 0.001
        one_blank = not number or not other_number
        rule = 0
        reason = ""
        if number and other_number == number and kind == "full":
            rule, reason = 1, "same receipt number"
        elif kind == "full" and same_total and gap <= 1 and one_blank:
            rule, reason = 2, "one day apart" if gap == 1 else "same day"
        elif kind == "full" and same_total and 1 < gap <= 7:
            rule, reason = 3, f"{gap} days apart"
        elif kind == "full" and gap <= 1 and not same_total:
            diff = abs(row["total"] - total)
            if 0 < diff <= _near_limit(row["total"], total) + 1e-9:
                rule, reason = 3, "close total"
        elif same_total and kind == "" and _merchant_unclear(store):
            rule, reason = 3, "unclear merchant"
        elif same_total and gap == 0 and kind == "partial":
            rule, reason = 3, "partial merchant"
        if rule:
            found.append({**row, "rule": rule, "reason": reason, "gap": gap})
    return found


def _covering_question(rows: list[dict], receipt: dict) -> dict | None:
    store = receipt.get("store", "")
    total = round(float(receipt["total"]), 2)
    day = _as_date(receipt.get("when"))
    if day is None:
        return None
    for row in rows:
        if row["status"] != "question":
            continue
        if merchant_match(row["merchant"], store) != "full":
            continue
        if abs(row["total"] - total) >= 0.001:
            continue
        if any(abs((listed - day).days) <= 7 for listed in _question_dates(row)):
            return row
    return None


def _why(candidate: dict) -> str:
    if candidate["reason"] == "close total":
        return "close total"
    if candidate["reason"] == "partial merchant":
        return "partial name"
    if candidate["reason"] == "unclear merchant":
        return "unclear merchant"
    gap = candidate["gap"]
    if gap == 0:
        return "same day"
    if gap == 1:
        return "one day apart"
    return f"{gap} days apart"


def _question_text(receipt: dict, candidates: list[dict]) -> str:
    store = short_store(receipt.get("store", ""))
    total = _money(receipt["total"])
    day = _as_date(receipt.get("when"))
    spoken = _spoken_date(day) if day else "that day"
    if len(candidates) == 1:
        other = _spoken_date(candidates[0]["date"])
        return f"Is the {store} receipt for ${total} on {spoken} the same purchase as the one on {other}?"
    pieces = [f"{_spoken_date(item['date'])} ({_why(item)})" for item in candidates]
    if len(pieces) == 2:
        listed = f"{pieces[0]}, or {pieces[1]}"
    else:
        listed = ", ".join(pieces[:-1]) + f", or {pieces[-1]}"
    dates = " or ".join(_spoken_date(item["date"]) for item in candidates)
    return (
        f"Is the {store} receipt for ${total} on {spoken} the same purchase as one of these: {listed}? "
        f"Reply with the date it matches, or no if it is new. "
        f"Which date does it match: {dates}? Reply no if it's new."
    )


def _source_label(receipt: dict) -> str:
    source = str(receipt.get("source") or "").strip()
    if source:
        return source
    ext = str(receipt.get("ext") or "")
    if ext == ".txt" or receipt.get("pasted"):
        return "text"
    return "photo"


def _file_label(saved_name: str, receipt: dict) -> str:
    source = _source_label(receipt)
    if source.startswith("email"):
        return "email"
    if source == "statement":
        return "statement"
    if saved_name.endswith(".txt") or source == "text":
        return "text"
    return "photo"


def apply_receipts(workbook_path, receipts: list[dict], receipts_dir, backups_dir, name_prefix: str = "", stamp: str | None = None) -> dict:
    """Write receipts into a local workbook.

    Matching follows the cloud book. Gather every candidate. Auto-attach only
    when there is exactly one rule 1 or rule 2 candidate. A near match asks and
    does not write an amount. A finished receipt key does not get a second
    amount. A Log row still marked ``pending`` is resumed. The workbook is not
    replaced when it is locked. A backup is taken before the file is replaced,
    and only when this call will change it.
    """
    workbook_path = Path(workbook_path)
    receipts_dir = Path(receipts_dir)
    backups_dir = Path(backups_dir)
    empty = {"locked": True, "wrote": [], "skipped": [], "resumed": [], "attached": [], "questions": [], "backup": ""}
    if is_locked(workbook_path):
        return empty
    wb = load_workbook(workbook_path)
    before = _grid_snapshot(wb)
    wrote: list[str] = []
    skipped: list[str] = []
    resumed: list[str] = []
    attached: list[str] = []
    questions: list[dict] = []
    touched: set[tuple] = set()
    seen: set[str] = set()
    reserved: set[tuple] = set()
    changed = False
    for receipt in receipts:
        items = _check_split(receipt)
        key = receipt_key(receipt.get("store", ""), receipt.get("number", ""), receipt.get("when"), receipt.get("total"))
        if key in seen:
            skipped.append(key)
            continue
        seen.add(key)
        outcome = _apply_one(wb, receipt, items, key, receipts_dir, name_prefix, reserved, touched)
        kind = outcome["action"]
        if kind == "skip":
            skipped.append(key)
        elif kind == "resume":
            resumed.append(key)
            changed = True
        elif kind == "attach":
            attached.append(key)
            changed = True
        elif kind == "question":
            questions.append({"key": key, "text": outcome["text"]})
            changed = True
        elif kind == "write":
            wrote.append(key)
            changed = True
        else:
            raise RuntimeError(f"unknown receipt action {kind}")

    if not changed:
        return {"locked": False, "wrote": [], "skipped": skipped, "resumed": [], "attached": [], "questions": [], "backup": ""}

    backup_path = backup_workbook(workbook_path, backups_dir, stamp=stamp)
    after = _grid_snapshot(wb)
    for coord, value in before.items():
        if coord in touched:
            continue
        if after.get(coord) != value:
            raise RuntimeError(f"unrelated cell changed: {coord}")
    _atomic_save(wb, workbook_path)
    return {
        "locked": False,
        "wrote": wrote,
        "skipped": skipped,
        "resumed": resumed,
        "attached": attached,
        "questions": questions,
        "backup": str(backup_path),
    }


def _apply_one(wb, receipt, items, key, receipts_dir, name_prefix, reserved, touched) -> dict:
    rows = _log_rows(wb)
    digest = sha256_bytes(_body_bytes(receipt)) if _body_bytes(receipt) else ""
    if _hash_seen(rows, digest):
        return {"action": "skip"}
    sheet, day, day_number = _sheet_for(receipt.get("when"))
    log_row = _find_log_row(wb, key)
    if log_row is not None:
        status = str(wb["Log"].cell(log_row, 14).value or "")
        if status == "pending":
            where = str(wb["Log"].cell(log_row, 5).value or "")
            _resume_pending(wb, log_row, receipt, items, key, sheet, day, day_number, where, receipts_dir, name_prefix, touched)
            return {"action": "resume"}
        if status == "question":
            text = _join_question(wb, log_row, receipt, receipts_dir, name_prefix)
            return {"action": "question", "text": text}
        if status != "superseded":
            _attach_to_log(wb, log_row, receipt, day, receipts_dir, name_prefix, touched)
            return {"action": "attach"}
    covering = _covering_question(rows, receipt)
    if covering is not None:
        text = _join_question(wb, covering["row"], receipt, receipts_dir, name_prefix)
        return {"action": "question", "text": text}
    candidates = _gather_candidates(_log_rows(wb), receipt)
    if len(candidates) == 1 and candidates[0]["rule"] in {1, 2}:
        _attach_to_log(wb, candidates[0]["row"], receipt, day, receipts_dir, name_prefix, touched)
        return {"action": "attach"}
    if candidates:
        text = _question_text(receipt, candidates)
        _write_question(wb, receipt, items, key, sheet, day, text, candidates, receipts_dir, name_prefix)
        return {"action": "question", "text": text}
    cells, fresh = _amount_cells(wb[sheet], day_number, items, reserved, store=receipt.get("store", ""))
    if not fresh:
        return {"action": "skip"}
    where = _where_text(sheet, cells)
    for _category, row, col, _amount in cells:
        touched.add((sheet, row, col))
    _write_new_receipt(wb, receipt, items, key, sheet, day, cells, where, receipts_dir, name_prefix)
    return {"action": "write"}


def _write_new_receipt(wb, receipt, items, key, sheet, day, cells, where, receipts_dir, name_prefix) -> None:
    ws = wb[sheet]
    label = short_store(receipt.get("store", ""))
    note = receipt.get("note") or ""
    for _category, row, col, amount in cells:
        current = ws.cell(row, col).value
        if isinstance(current, (int, float)) and abs(float(current) - amount) < 0.001:
            continue
        if current not in (None, ""):
            raise RuntimeError(f"refusing to overwrite {sheet}!{get_column_letter(col)}{row}")
        ws.cell(row, col - 1).value = f"{label} {note}".strip() if note else label
        ws.cell(row, col).value = amount
        ws.cell(row, col).number_format = "#,##0.00"
    log_row = _next_log_row(wb)
    receipt_row = _next_receipt_row(wb, sheet)
    _write_log(wb, log_row, receipt, where, day, sheet, receipt_row)
    _write_receipt_row(ws, receipt_row, receipt, items, key, where, day)
    saved_name = _store_body(receipt, receipts_dir, name_prefix)
    if saved_name:
        _mark_file(wb, log_row, ws, receipt_row, saved_name, _body_bytes(receipt), finished=True, cash=_is_cash(receipt))


def _resume_pending(wb, log_row, receipt, items, key, sheet, day, day_number, where, receipts_dir, name_prefix, touched) -> None:
    """Finish a pending row. Write an amount only into a cell that is still empty."""
    ws = wb[sheet]
    if where:
        for part in where.split("+"):
            cell_ref = part.split("!", 1)[-1]
            match = re.fullmatch(r"([A-Z]+)(\d+)", cell_ref)
            if not match:
                continue
            col = 0
            for letter in match.group(1):
                col = col * 26 + (ord(letter) - 64)
            row = int(match.group(2))
            current = ws.cell(row, col).value
            if current in (None, ""):
                ws.cell(row, col).value = round(float(receipt["total"]), 2) if "+" not in where else current
                touched.add((sheet, row, col))
            elif isinstance(current, (int, float)) and abs(float(current) - float(receipt["total"])) > 0.001 and "+" not in where:
                raise RuntimeError("pending cell holds a different amount")
    receipt_ws, receipt_row = _find_receipt_row(wb, key)
    if receipt_ws is None:
        if not where:
            cells, _fresh = _amount_cells(ws, day_number, items, touched, store=receipt.get("store", ""))
            where = _where_text(sheet, cells)
            wb["Log"].cell(log_row, 5).value = where
            for _category, row, col, amount in cells:
                current = ws.cell(row, col).value
                if isinstance(current, (int, float)) and abs(float(current) - amount) < 0.001:
                    continue
                if current not in (None, ""):
                    raise RuntimeError(f"refusing to overwrite {sheet}!{get_column_letter(col)}{row}")
                ws.cell(row, col).value = amount
                touched.add((sheet, row, col))
        receipt_row = _next_receipt_row(wb, sheet)
        _write_receipt_row(ws, receipt_row, receipt, items, key, where, day)
        wb["Log"].cell(log_row, 13).value = f"{sheet}!B{receipt_row}"
        receipt_ws = ws
    saved_name = _store_body(receipt, receipts_dir, name_prefix)
    if saved_name:
        _mark_file(wb, log_row, receipt_ws, receipt_row, saved_name, _body_bytes(receipt), finished=True, cash=_is_cash(receipt))


def _attach_to_log(wb, log_row: int, receipt: dict, new_day: date, receipts_dir, name_prefix: str, touched: set) -> None:
    """Add a file to an existing purchase. Do not add the amount. An earlier date keeps the amount."""
    log = wb["Log"]
    old_key = receipt_key(
        str(log.cell(log_row, 3).value or ""),
        str(log.cell(log_row, 9).value or ""),
        log.cell(log_row, 2).value,
        log.cell(log_row, 4).value or 0,
    )
    receipt_ws, receipt_row = _find_receipt_row(wb, old_key)
    old_day = _as_date(log.cell(log_row, 2).value)
    if old_day and new_day < old_day:
        _shift_amount_earlier(wb, log_row, receipt, new_day, touched)
        if receipt_ws is not None:
            hhmm = _as_hhmm(receipt.get("when"))
            receipt_ws.cell(receipt_row, 4).value = f"{new_day.isoformat()} {hhmm[:2]}:{hhmm[2:]}" if hhmm else new_day.isoformat()
            receipt_ws.cell(receipt_row, 79).value = receipt_key(
                str(log.cell(log_row, 3).value or ""),
                str(log.cell(log_row, 9).value or ""),
                new_day,
                log.cell(log_row, 4).value or 0,
            )
    source = _source_label(receipt)
    current = str(log.cell(log_row, 1).value or "")
    parts = [part.strip() for part in current.split(",")] if current else []
    if source and source not in parts:
        log.cell(log_row, 1).value = f"{current}, {source}" if current else source
    number = normalize_number(receipt.get("number", ""))
    if number and not str(log.cell(log_row, 9).value or "").strip():
        log.cell(log_row, 9).value = number
        if receipt_ws is not None and not str(receipt_ws.cell(receipt_row, 2).value or "").strip():
            receipt_ws.cell(receipt_row, 2).value = number
    saved_name = _store_body(receipt, receipts_dir, name_prefix)
    if saved_name:
        _append_file(wb, log_row, receipt_ws if receipt_ws is not None else wb[MONTHS[new_day.month - 1]], receipt_row, saved_name, _body_bytes(receipt), receipt)


def _shift_amount_earlier(wb, log_row: int, receipt: dict, new_day: date, touched: set) -> None:
    """Move the same amount onto the earlier day, on the same category row."""
    log = wb["Log"]
    where = str(log.cell(log_row, 5).value or "")
    if not where or "was =" in where:
        log.cell(log_row, 2).value = new_day.isoformat()
        return
    new_sheet = MONTHS[new_day.month - 1]
    _item_col, new_col = _day_columns(new_day.day)
    parts = []
    for piece in where.split("+"):
        piece = piece.strip()
        if "!" not in piece:
            parts.append(piece)
            continue
        sheet, ref = piece.split("!", 1)
        match = re.fullmatch(r"([A-Z]+)(\d+)", ref)
        if not match or sheet not in wb.sheetnames:
            parts.append(piece)
            continue
        col = 0
        for letter in match.group(1):
            col = col * 26 + (ord(letter) - 64)
        row = int(match.group(2))
        value = wb[sheet].cell(row, col).value
        target = wb[new_sheet].cell(row, new_col)
        if not isinstance(value, (int, float)):
            parts.append(piece)
            continue
        if target.value not in (None, "") and not (isinstance(target.value, (int, float)) and abs(float(target.value) - float(value)) < 0.001):
            parts.append(piece)
            continue
        target.value = float(value)
        target.number_format = "#,##0.00"
        if wb[new_sheet].cell(row, new_col - 1).value in (None, ""):
            wb[new_sheet].cell(row, new_col - 1).value = wb[sheet].cell(row, col - 1).value
        if not (sheet == new_sheet and col == new_col):
            wb[sheet].cell(row, col).value = None
            wb[sheet].cell(row, col - 1).value = None
        touched.add((sheet, row, col))
        touched.add((new_sheet, row, new_col))
        parts.append(f"{new_sheet}!{get_column_letter(new_col)}{row}")
    log.cell(log_row, 5).value = "+".join(parts)
    log.cell(log_row, 2).value = new_day.isoformat()
    hhmm = _as_hhmm(receipt.get("when"))
    log.cell(log_row, 10).value = f"{new_day.isoformat()} {hhmm[:2]}:{hhmm[2:]}" if hhmm else new_day.isoformat()


def _append_file(wb, log_row: int, ws, receipt_row: int | None, saved_name: str, data: bytes, receipt: dict) -> None:
    log = wb["Log"]
    label = _file_label(saved_name, receipt)
    current = str(log.cell(log_row, 6).value or "")
    count = current.count(": ") + 1 if current else 1
    pointer = f"{count} {label}: {saved_name}"
    log.cell(log_row, 6).value = f"{current}, {pointer}" if current else pointer
    digest = sha256_bytes(data) if data else ""
    if digest:
        hashes = str(log.cell(log_row, 15).value or "")
        log.cell(log_row, 15).value = f"{hashes}, {count} {digest}" if hashes else f"1 {digest}"
    if receipt_row is not None:
        existing = str(ws.cell(receipt_row, 27).value or "")
        if not existing:
            ws.cell(receipt_row, 27).value = label
        elif existing == label or existing.startswith(label + " +"):
            extra = 1
            match = re.search(r"\+(\d+)$", existing)
            if match:
                extra = int(match.group(1)) + 1
            ws.cell(receipt_row, 27).value = f"{label} +{extra}" if not match else f"{label} +{extra}"
        else:
            ws.cell(receipt_row, 27).value = f"{existing.split(' +')[0]} +1"


def _write_question(wb, receipt, items, key, sheet, day, text, candidates, receipts_dir, name_prefix) -> None:
    """One open question. No category amount is written."""
    log_row = _next_log_row(wb)
    log = wb["Log"]
    dates = ", ".join(item["date"].isoformat() for item in candidates if item["date"])
    source = _source_label(receipt)
    log.cell(log_row, 1).value = f"open question: candidates {dates}; {source}"
    log.cell(log_row, 2).value = day.isoformat()
    log.cell(log_row, 3).value = short_store(receipt.get("store", ""))
    log.cell(log_row, 4).value = round(float(receipt["total"]), 2)
    log.cell(log_row, 4).number_format = "#,##0.00"
    log.cell(log_row, 8).value = f"{day.isoformat()}; candidates {dates}"
    log.cell(log_row, 9).value = normalize_number(receipt.get("number", ""))
    hhmm = _as_hhmm(receipt.get("when"))
    log.cell(log_row, 10).value = f"{day.isoformat()} {hhmm[:2]}:{hhmm[2:]}" if hhmm else day.isoformat()
    log.cell(log_row, 11).value = "; ".join(f"{item['name']}: {float(item['price']):.2f}" for item in items)
    log.cell(log_row, 12).value = str(receipt.get("last4") or "")
    log.cell(log_row, 14).value = "question"
    ws = wb[sheet]
    receipt_row = _next_receipt_row(wb, sheet)
    _write_receipt_row(ws, receipt_row, receipt, items, key, "", day)
    ws.cell(receipt_row, 81).value = text
    ws.cell(receipt_row, 82).value = "question"
    saved_name = _store_body(receipt, receipts_dir, name_prefix)
    if saved_name:
        _append_file(wb, log_row, ws, receipt_row, saved_name, _body_bytes(receipt), receipt)


def _question_sheet_row(wb, log_row: int):
    merchant = wb["Log"].cell(log_row, 3).value
    total = wb["Log"].cell(log_row, 4).value or 0
    for ws, row in _iter_receipt_rows(wb):
        if ws.cell(row, 82).value != "question":
            continue
        if ws.cell(row, 3).value != merchant:
            continue
        amount = ws.cell(row, 6).value or 0
        if isinstance(amount, (int, float)) and abs(float(amount) - float(total)) < 0.001:
            return ws, row
    return None, None


def _join_question(wb, log_row: int, receipt: dict, receipts_dir, name_prefix: str) -> str:
    log = wb["Log"]
    day = _as_date(receipt.get("when"))
    notes = str(log.cell(log_row, 8).value or "")
    source = str(log.cell(log_row, 1).value or "")
    if day and day.isoformat() not in notes:
        if "candidates " in notes:
            notes = notes + f", {day.isoformat()}"
        else:
            notes = f"{day.isoformat()}; candidates {day.isoformat()}"
        log.cell(log_row, 8).value = notes
        prefix, sep, tail = source.partition(";")
        if "candidates " in prefix:
            log.cell(log_row, 1).value = prefix + f", {day.isoformat()}" + (sep + tail if sep else "")
    label = _source_label(receipt)
    source = str(log.cell(log_row, 1).value or "")
    if label and label not in source:
        log.cell(log_row, 1).value = f"{source}, {label}" if source else label
    ws, receipt_row = _question_sheet_row(wb, log_row)
    saved_name = _store_body(receipt, receipts_dir, name_prefix)
    if saved_name:
        _append_file(wb, log_row, ws if ws is not None else wb[MONTHS[0]], receipt_row, saved_name, _body_bytes(receipt), receipt)
    if ws is not None and receipt_row is not None and ws.cell(receipt_row, 81).value:
        return str(ws.cell(receipt_row, 81).value)
    return "Which date does it match? Reply no if it's new."


def _is_cash(receipt: dict) -> bool:
    return not str(receipt.get("last4") or "").strip()


def _write_log(wb, row: int, receipt: dict, where: str, day: date, sheet: str, receipt_row: int) -> None:
    log = wb["Log"]
    source = receipt.get("source") or "photo"
    log.cell(row, 1).value = source
    log.cell(row, 2).value = day.isoformat()
    log.cell(row, 3).value = short_store(receipt.get("store", ""))
    log.cell(row, 4).value = round(float(receipt["total"]), 2)
    log.cell(row, 4).number_format = "#,##0.00"
    log.cell(row, 5).value = where
    log.cell(row, 9).value = normalize_number(receipt.get("number", ""))
    hhmm = _as_hhmm(receipt.get("when"))
    log.cell(row, 10).value = f"{day.isoformat()} {hhmm[:2]}:{hhmm[2:]}" if hhmm else day.isoformat()
    items = _items_of(receipt)
    log.cell(row, 11).value = "; ".join(f"{item['name']}: {float(item['price']):.2f}" for item in items)
    log.cell(row, 12).value = str(receipt.get("last4") or "")
    log.cell(row, 13).value = f"{sheet}!B{receipt_row}"
    log.cell(row, 14).value = "pending"
    # Column P stays blank until a statement line exists. The receipt key lives on the Receipts row.


def _write_receipt_row(ws, row: int, receipt: dict, items: list[dict], key: str, where: str, day: date) -> None:
    ws.cell(row, 2).value = normalize_number(receipt.get("number", ""))
    ws.cell(row, 3).value = short_store(receipt.get("store", ""))
    hhmm = _as_hhmm(receipt.get("when"))
    ws.cell(row, 4).value = f"{day.isoformat()} {hhmm[:2]}:{hhmm[2:]}" if hhmm else day.isoformat()
    ws.cell(row, 6).value = round(float(receipt["total"]), 2)
    ws.cell(row, 6).number_format = "#,##0.00"
    ws.cell(row, 7).value = str(receipt.get("last4") or "") or "cash"
    ws.cell(row, 7).number_format = "@"
    for index, item in enumerate(items[:6]):
        start = 9 + index * 3
        ws.cell(row, start).value = item["name"]
        ws.cell(row, start + 1).value = round(float(item["price"]), 2)
        ws.cell(row, start + 1).number_format = "#,##0.00"
        ws.cell(row, start + 2).value = item["category"]
    ws.cell(row, 79).value = key
    ws.cell(row, 82).value = "pending"
    ws.cell(row, 83).value = where


def _store_body(receipt: dict, receipts_dir: Path, name_prefix: str) -> str:
    data = _body_bytes(receipt)
    if not data:
        return ""
    ext = receipt.get("ext") or ".txt"
    name = file_name(receipt.get("when"), receipt.get("store", ""), float(receipt["total"]), ext, prefix=name_prefix)
    return _place_file(receipts_dir, name, data)


def _mark_file(wb, log_row: int, ws, receipt_row: int, saved_name: str, data: bytes, finished: bool, cash: bool) -> None:
    label = "text" if saved_name.endswith(".txt") else "photo"
    pointer = f"1 {label}: {saved_name}"
    digest = sha256_bytes(data) if data else ""
    log = wb["Log"]
    log.cell(log_row, 6).value = pointer
    if digest:
        log.cell(log_row, 15).value = f"1 {digest}"
    if finished:
        log.cell(log_row, 14).value = "cash" if cash else "awaiting statement"
    ws.cell(receipt_row, 27).value = label
    ws.cell(receipt_row, 82).value = "cash" if cash else "pending"


def _has_statement_data(wb) -> bool:
    statements = wb["Statements"] if "Statements" in wb.sheetnames else None
    if statements is not None:
        for row in range(3, statements.max_row + 1):
            if statements.cell(row, 1).value:
                return True
    if "Cards" in wb.sheetnames:
        cards = wb["Cards"]
        for row in range(3, 31):
            if cards.cell(row, 3).value:
                return True
    for name in MONTHS:
        ws = wb[name]
        for row in range(49, 149):
            if ws.cell(row, 2).value or ws.cell(row, 6).value or ws.cell(row, 7).value:
                return True
    return False


def migrate_v21_settings(workbook_path) -> dict:
    """Bring a v2.0 book to v2.1 by filling settings cells only.

    Amounts, formulas, receipt rows, and transaction rows are not edited.
    Safe to run twice. A declined offer stays declined.
    """
    workbook_path = Path(workbook_path)
    if is_locked(workbook_path):
        return {"locked": True, "changed": False}
    wb = load_workbook(workbook_path)
    before = _grid_snapshot(wb)
    summary = _summary(wb)
    version = str(summary["C46"].value or "")
    if version not in {"v2.0", "v2.1"}:
        raise ValueError(f"v2.1 settings step only runs after v2.0 sections exist, not on {version!r}")
    changed = False
    if summary["C47"].value not in FEATURE_VALUES:
        summary["B47"] = "Statement features"
        summary["C47"] = "on" if _has_statement_data(wb) else "off"
        changed = True
    elif summary["B47"].value != "Statement features":
        summary["B47"] = "Statement features"
        changed = True
    if summary["C48"].value not in OFFER_VALUES:
        summary["B48"] = "Statement offer"
        summary["C48"] = "accepted" if summary["C47"].value == "on" else "not yet"
        changed = True
    elif summary["B48"].value != "Statement offer":
        summary["B48"] = "Statement offer"
        changed = True
    for row, label in ((49, "Last weekly report"), (50, "Last monthly report"), (51, "Last yearly report")):
        if summary.cell(row, 2).value != label:
            summary.cell(row, 2).value = label
            changed = True
    if summary["B52"].value != "Book started":
        summary["B52"] = "Book started"
        changed = True
    if not summary["C52"].value:
        # The day clock starts when this step runs. It does not use a receipt date.
        summary["C52"] = date.today().isoformat()
        changed = True
    banner = str(summary["B2"].value or "")
    if banner.startswith("Family Book v2.0"):
        summary["B2"] = "Family Book v2.1" + banner[len("Family Book v2.0"):]
        changed = True
    if version == "v2.0":
        summary["C46"] = "v2.1"
        changed = True
    after = _grid_snapshot(wb)
    if before != after:
        raise RuntimeError("v2.1 settings changed a number or a formula")
    if changed:
        _atomic_save(wb, workbook_path)
    return {"locked": False, "changed": changed, "version": _summary(load_workbook(workbook_path))["C46"].value if changed else version}


def mark_statement_added(workbook_path) -> dict:
    """The first imported statement turns statement features on. Amounts stay as they are."""
    workbook_path = Path(workbook_path)
    if is_locked(workbook_path):
        return {"locked": True}
    wb = load_workbook(workbook_path)
    before = _grid_snapshot(wb)
    summary = _summary(wb)
    summary["B47"] = "Statement features"
    summary["C47"] = "on"
    summary["B48"] = "Statement offer"
    summary["C48"] = "accepted"
    if before != _grid_snapshot(wb):
        raise RuntimeError("turning statement features on changed a number")
    _atomic_save(wb, workbook_path)
    return {"locked": False, "features": "on"}


def routine_gate(workbook_path) -> dict:
    """What the routines are allowed to do. Off until the first statement."""
    wb = load_workbook(workbook_path)
    settings = read_settings(wb)
    on = settings["features"] == "on"
    return {
        "features": settings["features"],
        "offer": settings["offer"],
        "may_remind": on,
        "may_nudge": on,
        "may_bank_income": on,
        "may_bank_fixed": on,
    }


def receipt_stats(workbook_path) -> dict:
    wb = load_workbook(workbook_path)
    count = 0
    first = None
    for ws, row in _iter_receipt_rows(wb):
        status = ws.cell(row, 82).value or ""
        if status in {"question", "superseded"}:
            continue
        count += 1
        day = _as_date(ws.cell(row, 4).value)
        if day and (first is None or day < first):
            first = day
    return {"count": count, "first": first}


def suggest_statement(features: str, offer: str, count: int, started: date | None, today: date) -> str:
    """One suggestion, or an empty string.

    ``started`` is Annual Summary ``C52`` (the day the book was set up), not
    the date on a receipt. A no, a yes, or an earlier ask returns nothing.
    """
    if features == "on":
        return ""
    if offer != "not yet":
        return ""
    days = (today - started).days if started else 0
    if count >= 10 or days >= 14:
        return SUGGESTION_TEXT
    return ""


def set_offer(workbook_path, offer: str) -> None:
    if offer not in OFFER_VALUES:
        raise ValueError(offer)
    workbook_path = Path(workbook_path)
    if is_locked(workbook_path):
        raise RuntimeError("workbook is locked")
    wb = load_workbook(workbook_path)
    before = _grid_snapshot(wb)
    summary = _summary(wb)
    summary["B48"] = "Statement offer"
    summary["C48"] = offer
    if before != _grid_snapshot(wb):
        raise RuntimeError("offer update changed a number")
    _atomic_save(wb, workbook_path)


def report_totals(workbook_path, start: date, end: date) -> dict:
    """Category sums for a date range, read from amount cells. Excel is not asked to calculate."""
    wb = load_workbook(workbook_path)
    year = _summary(wb)["F5"].value
    if not isinstance(year, int):
        raise ValueError("Annual Summary F5 has no year")
    totals = {name: 0.0 for name in CATEGORIES}
    for month_index, name in enumerate(MONTHS, 1):
        ws = wb[name]
        for day_number in range(1, 32):
            try:
                current = date(year, month_index, day_number)
            except ValueError:
                continue
            if current < start or current > end:
                continue
            _item_col, amount_col = _day_columns(day_number)
            for category, (row_start, row_end) in CATEGORIES.items():
                for row in range(row_start, row_end + 1):
                    value = ws.cell(row, amount_col).value
                    if isinstance(value, (int, float)):
                        totals[category] += float(value)
    return {key: round(value, 2) for key, value in totals.items() if value}


def advice(totals: dict) -> str:
    spent = {key: value for key, value in totals.items() if value}
    if not spent:
        return ""
    top = max(spent, key=spent.get)
    total = round(sum(spent.values()), 2)
    share = spent[top] / total if total else 0
    return f"{top} is the largest category, {spent[top]:.2f} of {total:.2f} ({share:.0%})."


def report_already_sent(workbook_path, kind: str, period: str) -> bool:
    wb = load_workbook(workbook_path)
    settings = read_settings(wb)
    return str(settings.get(kind) or "") == period


def record_report_sent(workbook_path, kind: str, period: str) -> None:
    cell = {"weekly": "C49", "monthly": "C50", "yearly": "C51"}[kind]
    workbook_path = Path(workbook_path)
    if is_locked(workbook_path):
        raise RuntimeError("workbook is locked")
    wb = load_workbook(workbook_path)
    before = _grid_snapshot(wb)
    _summary(wb)[cell] = period
    if before != _grid_snapshot(wb):
        raise RuntimeError("report marker changed a number")
    _atomic_save(wb, workbook_path)


def week_period(today: date) -> tuple[date, date, str]:
    """Monday-to-Sunday week ending on the most recent Sunday, including today.

    The id is that Sunday. A later run with the same id does not send again.
    """
    days_since_sunday = (today.weekday() + 1) % 7
    end = today - timedelta(days=days_since_sunday)
    start = end - timedelta(days=6)
    return start, end, end.isoformat()


def transactions_empty(workbook_path) -> bool:
    wb = load_workbook(workbook_path)
    for name in MONTHS:
        ws = wb[name]
        for row in range(49, 149):
            if any(ws.cell(row, col).value not in (None, "") for col in (2, 3, 4, 6, 7, 9, 10)):
                return False
    return True


def cards_empty(workbook_path) -> bool:
    wb = load_workbook(workbook_path)
    if "Cards" not in wb.sheetnames:
        return True
    cards = wb["Cards"]
    for row in range(3, 31):
        if cards.cell(row, 1).value or cards.cell(row, 3).value:
            return False
    return True
